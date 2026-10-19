# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Builds the app's local BLS asset from the official BLS 4.0 download.

Usage (from the repo root)::

    uv run tool/bls/bls_to_asset.py path/to/BLS_4_0_Daten_2025_DE.xlsx \
        assets/bls/bls_4_0.json.gz

Reads the data sheet of the Bundeslebensmittelschlüssel 4.0 (© Max
Rubner-Institut, CC BY 4.0) and writes a gzip-compressed JSON file::

    {"version": 1, "source": "...", "license": "...",
     "fields": ["energyKcal100", ...],
     "foods": [["C133000", "Hafer Flocken", "Oat flakes", [348, ...]], ...]}

Each food row holds its BLS code, German and English name and one value per
entry of ``fields``, per 100 g in the unit the app stores (see MAPPING); a
component BLS does not give for that food is ``null``, never 0.
"""

import gzip
import json
import sys
from pathlib import Path

ASSET_VERSION = 1
SOURCE = "Bundeslebensmittelschlüssel (BLS) 4.0, Max Rubner-Institut"
LICENSE = "CC BY 4.0, https://creativecommons.org/licenses/by/4.0/"

# App field -> (BLS component codes in priority order, unit factor to the
# app's unit). The first 23 match MealNutrimentsDBO and the backend mapping
# (trans fat has no BLS component), except that vitamin A takes only RAE:
# RE weighs carotenoids differently, so it is no stand-in. The rest are the
# extra micronutrients the app adds later (#50).
MAPPING = [
    ("energyKcal100", ["ENERCC"], 1),
    ("carbohydrates100", ["CHO"], 1),
    ("fat100", ["FAT"], 1),
    ("proteins100", ["PROT625"], 1),
    ("sugars100", ["SUGAR"], 1),
    ("saturatedFat100", ["FASAT"], 1),
    ("fiber100", ["FIBT"], 1),
    ("monounsaturatedFat100", ["FAMS"], 1),
    ("polyunsaturatedFat100", ["FAPU"], 1),
    ("cholesterol100", ["CHORL"], 1),
    ("sodium100", ["NA"], 1),
    ("potassium100", ["K"], 1),
    ("magnesium100", ["MG"], 1),
    ("calcium100", ["CA"], 1),
    ("iron100", ["FE"], 1),
    ("zinc100", ["ZN"], 1),
    ("phosphorus100", ["P"], 1),
    ("vitaminA100", ["VITAA"], 1),
    ("vitaminC100", ["VITC"], 1),
    ("vitaminD100", ["VITD"], 1),
    ("vitaminB6100", ["VITB6"], 0.001),
    ("vitaminB12100", ["VITB12"], 1),
    ("niacin100", ["NIA"], 1),
    ("vitaminE100", ["VITE"], 1),
    ("vitaminK100", ["VITK"], 1),
    ("thiamin100", ["THIA"], 1),
    ("riboflavin100", ["RIBF"], 1),
    ("pantothenicAcid100", ["PANTAC"], 1),
    ("biotin100", ["BIOT"], 1),
    ("folate100", ["FOL"], 1),
    ("iodine100", ["ID"], 1),
    ("copper100", ["CU"], 0.001),
    ("manganese100", ["MN"], 0.001),
    ("salt100", ["NACL"], 1),
    ("water100", ["WATER"], 1),
    ("alcohol100", ["ALC"], 1),
    ("omega3100", ["FAPUN3"], 1),
    ("omega6100", ["FAPUN6"], 1),
]


def parse_amount(value):
    """A BLS cell as a float: numbers, comma or dot strings; None for '-'."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text or text == "-":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def column_index(header):
    """Maps a BLS component code to its value column, e.g. 'ENERCC' -> 6."""
    columns = {}
    for index, title in enumerate(header):
        if title and "[" in str(title):
            columns[str(title).split(" ", 1)[0]] = index
    return columns


def convert_rows(header, rows):
    """The asset dict for the sheet given as a header and value rows."""
    columns = column_index(header)
    code_col = header.index("BLS Code")
    name_col = header.index("Lebensmittelbezeichnung")
    english_col = header.index("Food name")
    foods = []
    for row in rows:
        code = row[code_col]
        if not code:
            continue
        values = []
        for _, codes, factor in MAPPING:
            amount = None
            for component in codes:
                if component in columns:
                    amount = parse_amount(row[columns[component]])
                    if amount is not None:
                        break
            values.append(None if amount is None else round(amount * factor, 6))
        foods.append([str(code), row[name_col], row[english_col], values])
    return {
        "version": ASSET_VERSION,
        "source": SOURCE,
        "license": LICENSE,
        "fields": [field for field, _, _ in MAPPING],
        "foods": foods,
    }


def main(source, target):
    from openpyxl import load_workbook

    sheet = load_workbook(source, read_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    asset = convert_rows(header, rows)
    data = json.dumps(asset, ensure_ascii=False, separators=(",", ":"))
    Path(target).parent.mkdir(parents=True, exist_ok=True)
    with open(target, "wb") as raw, gzip.GzipFile(
        fileobj=raw, mode="wb", compresslevel=9, mtime=0
    ) as out:
        out.write(data.encode("utf-8"))
    print(f"{len(asset['foods'])} foods, {len(data) // 1024} KiB JSON -> {target}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
